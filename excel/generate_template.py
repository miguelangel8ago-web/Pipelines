"""Crea la PLANTILLA de Excel (estructura y estilos, sin datos).

Esto se corre UNA sola vez (o cuando cambie el diseño del reporte).
El script `fill_excel.py` nunca modifica la estructura: solo abre esta
plantilla y escribe valores en las celdas de datos.

Pestañas:
    - Resumen           KPIs generales del negocio
    - Zona              una fila por zona
    - Oficina           una fila por oficina (agrupada por zona)
    - SBU               una fila por unidad de negocio (SBU)
    - Detalle Empresas  encabezado; el fill agrega una fila por empresa
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.mock_data import SBUS, ZONAS

PLANTILLA_PATH = Path(__file__).parent / "plantilla_empresas.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
DATA_METRIC_COLS = [
    "# Empresas",
    "Empleados",
    "Ingresos Totales",
    "Costos Totales",
    "Utilidad Neta",
    "Margen Neto %",
]


def _estilizar_encabezado(ws, fila: int, columnas: list[str]):
    for col_idx, titulo in enumerate(columnas, start=1):
        celda = ws.cell(row=fila, column=col_idx, value=titulo)
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(18, len(titulo) + 4)


def _hoja_resumen(wb: Workbook):
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Reporte Ejecutivo de Empresas"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Generado el:"
    ws["B2"] = None  # lo llena fill_excel.py
    etiquetas = [
        "Total de Empresas",
        "Empleados Totales",
        "Ingresos Totales (MXN)",
        "Costos Totales (MXN)",
        "Utilidad Neta Total (MXN)",
        "Margen Neto Promedio (%)",
    ]
    fila_inicio = 4
    for i, etiqueta in enumerate(etiquetas):
        ws.cell(row=fila_inicio + i, column=1, value=etiqueta).font = Font(bold=True)
        # columna B queda vacía: aquí es donde escribe fill_excel.py
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22


def _hoja_por_grupo(wb: Workbook, nombre_hoja: str, etiqueta_col: str, filas_iniciales: list[dict]):
    ws = wb.create_sheet(nombre_hoja)
    columnas = [etiqueta_col] + DATA_METRIC_COLS
    _estilizar_encabezado(ws, 1, columnas)
    for i, datos_fila in enumerate(filas_iniciales, start=2):
        for col_idx, key in enumerate(columnas, start=1):
            if key in datos_fila:
                celda = ws.cell(row=i, column=col_idx, value=datos_fila[key])
                celda.border = THIN_BORDER
            else:
                ws.cell(row=i, column=col_idx).border = THIN_BORDER
    ws.freeze_panes = "A2"
    return ws


def _hoja_detalle(wb: Workbook):
    columnas = [
        "Empresa",
        "RFC",
        "Zona",
        "Oficina",
        "SBU",
        "Empleados",
        "Ingresos Últimos 12m",
        "Costos Últimos 12m",
        "Utilidad Neta 12m",
        "Margen Neto % 12m",
    ]
    ws = wb.create_sheet("Detalle Empresas")
    _estilizar_encabezado(ws, 1, columnas)
    ws.freeze_panes = "A2"
    return ws


def construir_plantilla() -> Path:
    wb = Workbook()
    _hoja_resumen(wb)

    filas_zona = [{"Zona": z} for z in ZONAS.keys()]
    _hoja_por_grupo(wb, "Zona", "Zona", filas_zona)

    filas_oficina = [
        {"Zona": zona, "Oficina": oficina}
        for zona, oficinas in ZONAS.items()
        for oficina in oficinas
    ]
    ws_oficina = wb.create_sheet("Oficina")
    columnas_oficina = ["Zona", "Oficina"] + DATA_METRIC_COLS
    _estilizar_encabezado(ws_oficina, 1, columnas_oficina)
    for i, fila in enumerate(filas_oficina, start=2):
        ws_oficina.cell(row=i, column=1, value=fila["Zona"]).border = THIN_BORDER
        ws_oficina.cell(row=i, column=2, value=fila["Oficina"]).border = THIN_BORDER
    ws_oficina.freeze_panes = "A2"

    filas_sbu = [{"SBU": s} for s in SBUS]
    _hoja_por_grupo(wb, "SBU", "SBU", filas_sbu)

    _hoja_detalle(wb)

    PLANTILLA_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(PLANTILLA_PATH)
    return PLANTILLA_PATH


if __name__ == "__main__":
    ruta = construir_plantilla()
    print(f"Plantilla creada en: {ruta}")
